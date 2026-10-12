"""
End-to-End Real Export Verification Test.
Executes the full pipeline:
PDF extraction -> Supplier Grouping & Bulk Review Resolution -> Excel Population & CSV Export.
Verifies all integrity constraints on a physical workbook and CSV output using synthetic fixtures.
"""
import os
import hashlib
import openpyxl
import pytest
from injaz.engine import CoreExtractionEngine
from injaz.excel_writer import ExcelPopulationEngine
from injaz.csv_writer import CsvPopulationEngine
from injaz.supplier_matcher import group_similar_suppliers
from injaz.models import UserDecision, DuplicateStatus


def test_real_end_to_end_export_pipeline(tmp_path, synthetic_pdf_path, synthetic_template_path):
    # 1. Verify fixture hashes before starting
    pdf_sha_before = hashlib.sha256(open(synthetic_pdf_path, "rb").read()).hexdigest()
    xlsx_sha_before = hashlib.sha256(open(synthetic_template_path, "rb").read()).hexdigest()

    # 2. Extract transactions using real engine
    engine = CoreExtractionEngine()
    result = engine.process_statement(synthetic_pdf_path)
    assert len(result.records) == 4

    # 3. Use Supplier Grouping to review & resolve font-spaced supplier artifacts
    groups = group_similar_suppliers(result.records)
    for grp in groups:
        canonical = grp.canonical_candidate or grp.suggested_label
        for r in grp.records:
            # Resolve supplier spacing concerns while preserving other reviews
            r.resolve_supplier_review(canonical, notes="تم الاعتماد والمراجعة الجماعية للمورد")

    # Resolve any remaining commercial inclusion decisions for testing
    for r in result.records:
        if r.requires_inclusion_review:
            r.user_decision = UserDecision.INCLUDED_BY_USER
        if r.duplicate_status == DuplicateStatus.POTENTIAL_DUPLICATE:
            r.resolve_duplicate_warning("تم التأكيد كمعاملة مستقلة")
        if r.requires_field_review and not r.field_review_resolved:
            # If invoice was single digit, confirm
            r.resolve_field_review("تأكيد رقم الفاتورة")

    # 4. Check eligible count
    eligible = [r for r in result.records if r.is_export_eligible]
    assert len(eligible) == 4

    # 5. Export to Excel in a dedicated output path in tmp_path
    out_dir = tmp_path / "exports"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_xlsx = str(out_dir / "e2e_verified_export.xlsx")

    excel_writer = ExcelPopulationEngine()
    summary = excel_writer.export_transactions(
        template_path=synthetic_template_path,
        output_path=out_xlsx,
        records=result.records,
        overwrite=True
    )

    # 6. Verify ExportSummary
    assert summary.total_records_submitted == 4
    assert summary.records_rejected_ineligible == 0
    assert summary.records_exported + summary.records_skipped_already_present == 4
    assert os.path.exists(out_xlsx)
    assert os.path.getsize(out_xlsx) > 0

    # 7. Reopen output workbook and verify structure
    wb = openpyxl.load_workbook(out_xlsx, data_only=False)
    assert "تحليلي مشتريات" in wb.sheetnames
    ws = wb["تحليلي مشتريات"]

    # Table headers are on row 2
    assert ws.cell(2, 1).value == "م"
    assert ws.cell(2, 2).value == "التاريخ"
    assert ws.cell(2, 3).value == "المشروع"
    assert ws.cell(2, 4).value == "رقم الفاتورة"
    assert ws.cell(2, 5).value == "المورد"

    # Verify formulas in Value (F) and Total (H)
    for r_idx in range(3, min(ws.max_row, 30)):
        f_val = ws.cell(r_idx, 6).value
        f_tot = ws.cell(r_idx, 8).value
        # Formulas should start with '='
        if f_val:
            assert str(f_val).startswith("=")
        if f_tot:
            assert str(f_tot).startswith("=")

    # 8. Export to CSV Fallback
    out_csv = str(out_dir / "e2e_verified_export.csv")
    csv_writer = CsvPopulationEngine()
    csv_summary = csv_writer.export_transactions(
        output_path=out_csv,
        records=result.records,
        overwrite=True
    )

    assert csv_summary.records_exported == 4
    assert csv_summary.records_rejected_ineligible == 0
    assert os.path.exists(out_csv)
    assert os.path.getsize(out_csv) > 0

    # Verify CSV content
    with open(out_csv, "r", newline="", encoding="utf-8-sig") as f:
        import csv
        reader = list(csv.reader(f))
    assert reader[0] == ["م", "التاريخ", "المشروع", "رقم الفاتورة", "المورد", "الضريبة"]
    assert len(reader) == 5  # 1 header + 4 data rows

    # 9. Verify fixture hashes remain completely untouched
    assert hashlib.sha256(open(synthetic_pdf_path, "rb").read()).hexdigest() == pdf_sha_before
    assert hashlib.sha256(open(synthetic_template_path, "rb").read()).hexdigest() == xlsx_sha_before

