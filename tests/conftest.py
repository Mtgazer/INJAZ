"""
Pytest fixtures and synthetic data generators for INJAZ test suite.
Provides 100% synthetic, privacy-safe PDF statements and Excel templates.
"""
import os
import pytest
import pymupdf
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font

from injaz.models import (
    TransactionRecord,
    ExtractionStatus,
    InclusionEligibility,
    UserDecision,
    DuplicateStatus,
)


@pytest.fixture(scope="session")
def fixtures_dir():
    """Directory containing reusable synthetic fixtures."""
    d = os.path.join(os.path.dirname(__file__), "fixtures")
    os.makedirs(d, exist_ok=True)
    return d


@pytest.fixture(autouse=True)
def clean_qt_windows():
    """Closes all top-level Qt widgets and processes pending events after each test."""
    yield
    try:
        import gc
        from PySide6.QtWidgets import QApplication
        app = QApplication.instance()
        if app:
            for widget in app.topLevelWidgets():
                widget.close()
                widget.deleteLater()
            app.processEvents()
            app.sendPostedEvents()
        gc.collect()
    except Exception:
        pass





def _create_synthetic_template(dest_path: str):
    """Builds synthetic template completely from scratch with zero private data."""
    import datetime
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    wb.properties.creator = "INJAZ Test Suite"
    wb.properties.lastModifiedBy = "INJAZ Test Suite"
    wb.properties.title = "Synthetic Demonstration Template"
    wb.properties.subject = "INJAZ Test Fixture"
    wb.properties.description = "Synthetic test template built 100% from scratch with fictitious data."
    
    ws = wb.create_sheet(title="تحليلي مشتريات")
    ws.sheet_view.rightToLeft = True
    
    ws.cell(1, 1).value = "تحليلي مشتريات — نموذج تجريبي افتراضي"
    ws.cell(1, 1).font = Font(name="Simplified Arabic", size=14, bold=True)
    
    headers = [
        "م", "التاريخ", "المشروع", "رقم الفاتورة", "المورد",
        "القيمة", "الضريبة", "الإجمالي", "ملاحظات", "تصنيف", "حساب", "مركز"
    ]
    for c_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(2, c_idx, h_text)
        cell.font = Font(name="Simplified Arabic", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="16324F", end_color="16324F", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    arabic_font = Font(name="Simplified Arabic", size=11)
    thin_side = Side(style="thin", color="CBD5E1")
    grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    fake_projects = [
        "مشروع صيل 10", "مشروع صيل 11", "مشروع صيل 12", "مشروع صيل 14",
        "مشروع برج النرجس", "مشروع واحة الياسمين", "مشروع مجمع الفيحاء"
    ]
    fake_suppliers = [
        "شركة الأمل للتوريدات الإنشائية", "مؤسسة البناء والتعمير الحديثة",
        "شركة التقنية الهندسية ذ.م.م", "مؤسسة الرواد للتجارة والمقاولات",
        "شركة المواد المتقدمة المحدودة", "مؤسسة الأفق للمقاولات العامة",
        "شركة الرواف للتجارة والمقاولات", "مصنع الأسمنت الوطني الحديث"
    ]
    
    base_date = datetime.date(2026, 1, 1)
    for r_idx in range(3, 192):
        seq = r_idx - 2
        dt_val = (base_date + datetime.timedelta(days=seq)).strftime("%Y-%m-%d")
        proj_val = fake_projects[seq % len(fake_projects)]
        inv_val = 10000 + seq
        supp_val = fake_suppliers[seq % len(fake_suppliers)]
        tax_val = round(15.0 + (seq * 0.75), 2)
        
        ws.cell(r_idx, 1, seq).alignment = Alignment(horizontal="center")
        c2 = ws.cell(r_idx, 2, dt_val)
        c2.number_format = "yyyy-mm-dd"
        c2.alignment = Alignment(horizontal="center")
        ws.cell(r_idx, 3, proj_val).font = arabic_font
        ws.cell(r_idx, 4, inv_val).alignment = Alignment(horizontal="center")
        ws.cell(r_idx, 5, supp_val).font = arabic_font
        ws.cell(r_idx, 6, "=Table6[[#This Row],[الضريبة]]/0.15").number_format = "#,##0.00"
        ws.cell(r_idx, 7, tax_val).number_format = "#,##0.00"
        ws.cell(r_idx, 8, "=Table6[[#This Row],[القيمة]]+Table6[[#This Row],[الضريبة]]").number_format = "#,##0.00"
        
        for c in range(1, 13):
            ws.cell(r_idx, c).border = grid_border
            
    for r_idx in range(192, 375):
        seq = r_idx - 2
        ws.cell(r_idx, 1, seq).alignment = Alignment(horizontal="center")
        ws.cell(r_idx, 2, None).number_format = "yyyy-mm-dd"
        ws.cell(r_idx, 3, None).font = arabic_font
        ws.cell(r_idx, 4, None)
        ws.cell(r_idx, 5, None).font = arabic_font
        ws.cell(r_idx, 6, "=Table6[[#This Row],[الضريبة]]/0.15").number_format = "#,##0.00"
        ws.cell(r_idx, 7, None).number_format = "#,##0.00"
        ws.cell(r_idx, 8, "=Table6[[#This Row],[القيمة]]+Table6[[#This Row],[الضريبة]]").number_format = "#,##0.00"
        for c in range(1, 13):
            ws.cell(r_idx, c).border = grid_border
            
    tbl = Table(displayName="Table6", ref="A2:L374")
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    for i, h in enumerate(headers, 1):
        tbl.tableColumns.append(TableColumn(id=i, name=h))
    ws.add_table(tbl)
    
    r_sum = 375
    ws.cell(r_sum, 1, "الإجمالي").font = Font(name="Simplified Arabic", size=11, bold=True)
    ws.cell(r_sum, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(r_sum, 6, "=SUBTOTAL(109,Table6[القيمة])").number_format = "#,##0.00"
    ws.cell(r_sum, 7, "=SUBTOTAL(109,Table6[الضريبة])").number_format = "#,##0.00"
    ws.cell(r_sum, 8, "=SUBTOTAL(109,Table6[الإجمالي])").number_format = "#,##0.00"
    for c in range(1, 13):
        ws.cell(r_sum, c).border = grid_border
    ws.merge_cells("A375:E375")
    ws.merge_cells("I375:L375")
    
    ws_sheet1 = wb.create_sheet(title="Sheet1")
    ws_sheet1.cell(1, 1, "بيانات الشركة النموذجية")
    ws_sheet1.cell(4, 2, "شركة النماذج الافتراضية")
    ws_sheet1.cell(4, 3, "3999999999")
    
    ws_summary = wb.create_sheet(title="ملخص")
    ws_summary.cell(1, 1, "المشروع")
    ws_summary.cell(1, 2, "إجمالي الضريبة")
    for idx, p in enumerate(fake_projects, 2):
        ws_summary.cell(idx, 1, p)
        ws_summary.cell(idx, 2, 1000.0 * idx)
        
    os.makedirs(os.path.dirname(dest_path), exist_ok=True)
    wb.save(dest_path)
    wb.close()


@pytest.fixture(scope="session")
def synthetic_template_path(fixtures_dir):
    """
    Returns path to a synthetic accounting template workbook, generating from scratch if needed.
    """
    template_path = os.path.join(fixtures_dir, "synthetic_template.xlsx")
    if not os.path.exists(template_path):
        _create_synthetic_template(template_path)
    return template_path


@pytest.fixture(scope="session")
def synthetic_pdf_path(fixtures_dir):
    """
    Generates a synthetic financial statement PDF using PyMuPDF.
    Contains multi-column financial transaction rows with dates, numbers,
    descriptions, invoice numbers, supplier names, rule exclusions,
    and duplicate candidates.
    """
    pdf_path = os.path.join(fixtures_dir, "synthetic_statement.pdf")
    
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    
    # Statement header (above y=260)
    page.insert_text((180, 50), "Statement of Account - Synthetic Demonstration")
    page.insert_text((180, 70), "Account No: 100200300400")
    
    # 4 Structured Transaction rows
    rows = [
        # (y, date, type, supp, inv, desc, ref, debit, credit, bal, proj)
        (280, "2026-05-01", "Entry", "Al-Amal Co", "101", "Building Materials", "REF001", "150.00", "0.00", "150.00", "Proj 14"),
        (315, "2026-05-02", "Entry", "Utility Co", "102", "Office Supplies", "REF002", "300.00", "0.00", "450.00", "HQ Office"),
        (350, "2026-05-03", "Entry", "Modern Paint", "103", "Finishing Work", "REF003", "500.00", "0.00", "950.00", "Proj 16"),
        (385, "2026-05-03", "Entry", "Modern Paint", "103", "Finishing Work", "REF004", "500.00", "0.00", "1450.00", "Proj 16"),
    ]
    
    for y, dt, tp, supp, inv, desc, ref, deb, cred, bal, proj in rows:
        page.insert_text((505, y), dt)
        page.insert_text((475, y), tp)
        page.insert_text((380, y), supp)
        page.insert_text((320, y), inv)
        page.insert_text((260, y), desc)
        page.insert_text((220, y), ref)
        page.insert_text((180, y), deb)
        page.insert_text((140, y), cred)
        page.insert_text((90, y), bal)
        page.insert_text((55, y), proj)
        
    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def sample_synthetic_records():
    """Generates a rich batch of synthetic TransactionRecords covering all eligibility states."""
    records = []
    
    # 1. Validated and Eligible
    for i in range(1, 11):
        records.append(TransactionRecord(
            page_number=1,
            record_index=i,
            date=f"2026-06-{i:02d}",
            transaction_type="قيد يدوي",
            raw_description=f"توريد مواد بناء مشروع صيل 14 فاتورة 100{i} شركة الرواف للتجارة",
            reference=f"REF-{i:04d}",
            debit=100.0 * i,
            credit=0.0,
            balance=str(100.0 * i),
            project="صيل 14",
            invoice_number=f"100{i}",
            supplier_name="شركة الرواف للتجارة",
            tax_amount=15.0 * i,
            extraction_status=ExtractionStatus.VALIDATED,
            inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
            user_decision=UserDecision.NOT_APPLICABLE,
            duplicate_status=DuplicateStatus.UNIQUE,
        ))
        
    # 2. Excluded by rule (electricity / rent)
    records.append(TransactionRecord(
        page_number=1,
        record_index=11,
        date="2026-06-11",
        transaction_type="قيد يدوي",
        raw_description="سداد فاتورة كهرباء المركز الرئيسي",
        reference="REF-0011",
        debit=350.0,
        credit=0.0,
        balance="350.0",
        project="المركز الرئيسي",
        invoice_number="9901",
        supplier_name="شركة الكهرباء",
        tax_amount=52.5,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.EXCLUDED_BY_RULE,
        exclusion_reasons=["فاتورة كهرباء المركز الرئيسي"],
        user_decision=UserDecision.EXCLUDED_BY_RULE,
        duplicate_status=DuplicateStatus.UNIQUE,
    ))
    
    # 3. Requires Field Review (single-digit invoice)
    records.append(TransactionRecord(
        page_number=1,
        record_index=12,
        date="2026-06-12",
        transaction_type="قيد يدوي",
        raw_description="شراء أدوات سباكة فاتورة 4 مؤسسة الأمل",
        reference="REF-0012",
        debit=120.0,
        credit=0.0,
        balance="120.0",
        project="صيل 14",
        invoice_number="4",
        supplier_name="مؤسسة الأمل",
        tax_amount=18.0,
        extraction_status=ExtractionStatus.MANUAL_REVIEW_REQUIRED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        review_reasons=["single-digit invoice number requires verification"],
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE,
    ))
    
    # 4. Duplicate candidates
    records.append(TransactionRecord(
        page_number=1,
        record_index=13,
        date="2026-06-13",
        transaction_type="قيد يدوي",
        raw_description="دهانات فاتورة 55 شركة الألوان",
        reference="REF-0013",
        debit=200.0,
        credit=0.0,
        balance="200.0",
        project="صيل 14",
        invoice_number="55",
        supplier_name="شركة الألوان",
        tax_amount=30.0,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.POTENTIAL_DUPLICATE,
        duplicate_group_id="DUP-GRP-01",
    ))
    
    return records
