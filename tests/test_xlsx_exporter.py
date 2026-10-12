"""
Unit and Integration Tests for INJAZ (إنجاز) New Excel Exporter (xlsx_exporter.py).
Validates generation of brand-new .xlsx workbooks from scratch, strict fail-closed
audit rules, formatting, RTL, parity with CSV, and absence of template dependencies.
"""
import sys
import os
import hashlib
import datetime
from typing import List

import pytest
import openpyxl
from PySide6.QtWidgets import QApplication, QMessageBox, QDialog
from PySide6.QtCore import Qt


@pytest.fixture(scope="session")
def qapp():
    """Initializes a shared QApplication in offscreen mode with Arabic RTL layout."""
    app = QApplication.instance()
    if app is None:
        app = QApplication(sys.argv)
    app.setLayoutDirection(Qt.RightToLeft)
    return app


@pytest.fixture(autouse=True)
def mock_message_boxes(monkeypatch):
    """Prevents modal dialogs from blocking headless tests."""
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, "exec", lambda self: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.Ok)
    monkeypatch.setattr(QDialog, "exec", lambda self: QDialog.Accepted)
    monkeypatch.setattr(QDialog, "exec_", lambda self: QDialog.Accepted)

from injaz.models import (
    TransactionRecord,
    ExtractionStatus,
    InclusionEligibility,
    UserDecision,
    DuplicateStatus,
    ExclusionCategory,
)
from injaz.xlsx_exporter import (
    XlsxExporter,
    XlsxExportSummary,
    XlsxExportError,
    NoEligibleRecordsError,
    IneligibleXlsxExportError,
    XlsxProtectionError,
    XlsxValidationError,
)
from injaz.csv_writer import CsvPopulationEngine
from injaz.ui.pages.export_page import ExportPage


def create_synthetic_record(
    index: int = 1,
    date: str = "2026-05-12",
    project: str = "مشروع الرياض التجاري",
    invoice_number: str = "INV-100200",
    supplier_name: str = "شركة الفنار للمقاولات",
    tax_amount: float = 1500.0,
    extraction_status: ExtractionStatus = ExtractionStatus.VALIDATED,
    inclusion_eligibility: InclusionEligibility = InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
    user_decision: UserDecision = UserDecision.NOT_APPLICABLE,
    duplicate_status: DuplicateStatus = DuplicateStatus.UNIQUE,
    field_review_resolved: bool = True,
    duplicate_resolved: bool = True,
) -> TransactionRecord:
    rec = TransactionRecord(
        page_number=1,
        record_index=index,
        date=date,
        transaction_type="قيد يدوي",
        raw_description=f"فاتورة {invoice_number} من {supplier_name}",
        reference="REF-001",
        debit=tax_amount if tax_amount >= 0 else 0.0,
        credit=-tax_amount if tax_amount < 0 else 0.0,
        balance="10000.00",
        project=project,
        invoice_number=invoice_number,
        supplier_name=supplier_name,
        tax_amount=tax_amount,
        extraction_status=extraction_status,
        inclusion_eligibility=inclusion_eligibility,
        user_decision=user_decision,
        duplicate_status=duplicate_status,
        field_review_resolved=field_review_resolved,
        duplicate_resolved=duplicate_resolved,
    )
    return rec


# =========================================================================
# 1. Successful creation of a new .xlsx workbook
# =========================================================================
def test_1_successful_creation_of_new_xlsx_workbook(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_out.xlsx")
    records = [create_synthetic_record(1), create_synthetic_record(2)]

    summary = exporter.export_transactions(output_path, records)

    assert os.path.exists(output_path)
    assert os.path.getsize(output_path) > 0
    assert isinstance(summary, XlsxExportSummary)
    assert summary.records_exported == 2
    assert summary.total_records_submitted == 2
    assert summary.records_rejected_ineligible == 0


# =========================================================================
# 2. Exact worksheet name and six-column header order
# =========================================================================
def test_2_exact_worksheet_name_and_six_column_header_order(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_headers.xlsx")
    records = [create_synthetic_record(1)]

    exporter.export_transactions(output_path, records)

    wb = openpyxl.load_workbook(output_path)
    assert wb.sheetnames == ["المعاملات"]
    ws = wb["المعاملات"]

    expected_headers = ["م", "التاريخ", "المشروع", "رقم الفاتورة", "المورد", "الضريبة"]
    actual_headers = [ws.cell(row=1, column=c).value for c in range(1, 7)]
    assert actual_headers == expected_headers
    wb.close()


# =========================================================================
# 3. Correct sequential numbering (1..N)
# =========================================================================
def test_3_correct_sequential_numbering(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_seq.xlsx")
    # Records with non-consecutive internal indices
    records = [
        create_synthetic_record(15),
        create_synthetic_record(42),
        create_synthetic_record(99),
    ]

    exporter.export_transactions(output_path, records)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["المعاملات"]
    seq_col_values = [ws.cell(row=r, column=1).value for r in range(2, 5)]
    assert seq_col_values == [1, 2, 3]
    wb.close()


# =========================================================================
# 4. Correct dates and numeric tax values
# =========================================================================
def test_4_correct_dates_and_numeric_tax_values(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_types.xlsx")
    records = [create_synthetic_record(1, date="2026-10-12", tax_amount=1234.56)]

    exporter.export_transactions(output_path, records)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["المعاملات"]

    # Date cell (B2)
    cell_date = ws.cell(row=2, column=2)
    assert isinstance(cell_date.value, (datetime.date, datetime.datetime))
    assert cell_date.number_format.lower() in ["yyyy-mm-dd", "yyyy\\-mm\\-dd"]

    # Tax cell (F2)
    cell_tax = ws.cell(row=2, column=6)
    assert isinstance(cell_tax.value, (int, float))
    assert cell_tax.value == pytest.approx(1234.56)
    assert cell_tax.number_format == "#,##0.00"
    wb.close()


# =========================================================================
# 5. Positive and negative tax signs
# =========================================================================
def test_5_positive_and_negative_tax_signs(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_signs.xlsx")
    rec_pos = create_synthetic_record(1, tax_amount=850.25)
    rec_neg = create_synthetic_record(2, tax_amount=-420.50)

    exporter.export_transactions(output_path, [rec_pos, rec_neg])

    wb = openpyxl.load_workbook(output_path)
    ws = wb["المعاملات"]
    assert ws.cell(row=2, column=6).value == pytest.approx(850.25)
    assert ws.cell(row=3, column=6).value == pytest.approx(-420.50)
    wb.close()


# =========================================================================
# 6. Arabic supplier and project names
# =========================================================================
def test_6_arabic_supplier_and_project_names(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_arabic.xlsx")
    ar_proj = "مشروع برج الأندلس السكني"
    ar_supp = "شركة التوريدات الإنشائية المتطورة ذ.م.م"
    rec = create_synthetic_record(1, project=ar_proj, supplier_name=ar_supp)

    exporter.export_transactions(output_path, [rec])

    wb = openpyxl.load_workbook(output_path)
    ws = wb["المعاملات"]
    assert ws.cell(row=2, column=3).value == ar_proj
    assert ws.cell(row=2, column=5).value == ar_supp
    wb.close()


# =========================================================================
# 7. Bulk supplier corrections appearing in exported data
# =========================================================================
def test_7_bulk_supplier_corrections_appearing_in_exported_data(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_bulk_edit.xlsx")
    r1 = create_synthetic_record(1, supplier_name="مؤسسة النور")
    r2 = create_synthetic_record(2, supplier_name="مؤسسة  النور ")
    r3 = create_synthetic_record(3, supplier_name="مؤسسة النور للتجارة")

    # Apply approved bulk correction
    canonical = "شركة النور للمقاولات المعتمدة"
    for r in (r1, r2, r3):
        r.resolve_supplier_review(canonical)

    exporter.export_transactions(output_path, [r1, r2, r3])

    wb = openpyxl.load_workbook(output_path)
    ws = wb["المعاملات"]
    for row in range(2, 5):
        assert ws.cell(row=row, column=5).value == canonical
    wb.close()


# =========================================================================
# 8. Rejection of unresolved field-review records
# =========================================================================
def test_8_rejection_of_unresolved_field_review_records(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_field_rev.xlsx")
    valid_rec = create_synthetic_record(1)
    
    unresolved_rec = create_synthetic_record(
        2,
        extraction_status=ExtractionStatus.MANUAL_REVIEW,
        field_review_resolved=False
    )
    unresolved_rec.add_review_reason("مورد يحتاج تأكيد")

    summary = exporter.export_transactions(output_path, [valid_rec, unresolved_rec])

    assert summary.records_exported == 1
    assert summary.records_rejected_ineligible == 1
    assert summary.rejection_details[0]["record_index"] == 2

    wb = openpyxl.load_workbook(output_path)
    ws = wb["المعاملات"]
    assert ws.max_row == 2  # 1 header + 1 eligible data row
    wb.close()


# =========================================================================
# 9. Rejection of unresolved inclusion decisions
# =========================================================================
def test_9_rejection_of_unresolved_inclusion_decisions(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_inc_rev.xlsx")

    # Requires review and not yet approved
    unapproved_rec = create_synthetic_record(
        1,
        inclusion_eligibility=InclusionEligibility.REQUIRES_USER_REVIEW,
        user_decision=UserDecision.PENDING
    )

    with pytest.raises(NoEligibleRecordsError):
        exporter.export_transactions(output_path, [unapproved_rec])

    # Now approve it
    unapproved_rec.user_decision = UserDecision.INCLUDED_BY_USER
    summary = exporter.export_transactions(output_path, [unapproved_rec])
    assert summary.records_exported == 1


# =========================================================================
# 10. Rejection of unresolved duplicate warnings
# =========================================================================
def test_10_rejection_of_unresolved_duplicate_warnings(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_dup_warn.xlsx")

    dup_rec = create_synthetic_record(
        1,
        duplicate_status=DuplicateStatus.POTENTIAL_DUPLICATE,
        duplicate_resolved=False
    )

    with pytest.raises(NoEligibleRecordsError):
        exporter.export_transactions(output_path, [dup_rec])

    # Resolve duplicate warning
    dup_rec.resolve_duplicate_warning("Approved duplicate by auditor")
    summary = exporter.export_transactions(output_path, [dup_rec])
    assert summary.records_exported == 1


# =========================================================================
# 11. Exclusion of records rejected by business rules
# =========================================================================
def test_11_exclusion_of_records_rejected_by_business_rules(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_rule_excl.xlsx")

    excl_rec = create_synthetic_record(
        1,
        inclusion_eligibility=InclusionEligibility.EXCLUDED_BY_RULE
    )
    excl_rec.exclusion_category = ExclusionCategory.ELECTRICITY_BILL

    with pytest.raises(NoEligibleRecordsError):
        exporter.export_transactions(output_path, [excl_rec])


# =========================================================================
# 12. No empty-success export when no eligible records exist
# =========================================================================
def test_12_no_empty_success_export_when_no_eligible_records_exist(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_no_empty.xlsx")

    # Empty list
    with pytest.raises(NoEligibleRecordsError):
        exporter.export_transactions(output_path, [])
    assert not os.path.exists(output_path)

    # Ineligible record only
    bad_rec = create_synthetic_record(1)
    bad_rec.user_decision = UserDecision.EXCLUDED_BY_USER

    with pytest.raises(NoEligibleRecordsError):
        exporter.export_transactions(output_path, [bad_rec])
    assert not os.path.exists(output_path)


# =========================================================================
# 13. Existing output path is not overwritten silently
# =========================================================================
def test_13_existing_output_path_is_not_overwritten_silently(tmp_path):
    exporter = XlsxExporter()
    existing_file = tmp_path / "existing.xlsx"
    existing_file.write_text("original content")

    records = [create_synthetic_record(1)]

    # Without overwrite=True: must raise FileExistsError
    with pytest.raises(FileExistsError):
        exporter.export_transactions(str(existing_file), records, overwrite=False)

    assert existing_file.read_text() == "original content"

    # With overwrite=True: replaces cleanly
    summary = exporter.export_transactions(str(existing_file), records, overwrite=True)
    assert summary.records_exported == 1
    assert os.path.getsize(str(existing_file)) > 50


# =========================================================================
# 14. Temporary output is reopened and validated
# =========================================================================
def test_14_temporary_output_is_reopened_and_validated(tmp_path, monkeypatch):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "test_val_fail.xlsx")
    records = [create_synthetic_record(1)]

    # Mock _build_workbook to create corrupt workbook with wrong headers
    def corrupt_build(path, recs):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "المعاملات"
        ws.append(["wrong", "headers"])
        wb.save(path)

    monkeypatch.setattr(exporter, "_build_workbook", corrupt_build)

    with pytest.raises(XlsxValidationError):
        exporter.export_transactions(output_path, records)

    # Target output file must NOT have been moved/finalized
    assert not os.path.exists(output_path)


# =========================================================================
# 15. CSV and XLSX contain the same approved dataset
# =========================================================================
def test_15_csv_and_xlsx_contain_the_same_approved_dataset(tmp_path):
    xlsx_exporter = XlsxExporter()
    csv_engine = CsvPopulationEngine()

    xlsx_path = str(tmp_path / "compare.xlsx")
    csv_path = str(tmp_path / "compare.csv")

    records = [
        create_synthetic_record(1, date="2026-01-10", project="مشروع أ", invoice_number="INV-1", supplier_name="مورد أ", tax_amount=100.50),
        create_synthetic_record(2, date="2026-02-15", project="مشروع ب", invoice_number="INV-2", supplier_name="مورد ب", tax_amount=-50.25),
        create_synthetic_record(3, date="2026-03-20", project="مشروع ج", invoice_number="INV-3", supplier_name="مورد ج", tax_amount=300.00),
    ]

    xlsx_summary = xlsx_exporter.export_transactions(xlsx_path, records)
    csv_summary = csv_engine.export_transactions(csv_path, records)

    assert xlsx_summary.records_exported == csv_summary.records_exported == 3

    # Load XLSX
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["المعاملات"]

    # Load CSV
    import csv
    with open(csv_path, "r", encoding="utf-8-sig") as f:
        csv_rows = list(csv.reader(f))

    # Compare headers
    xlsx_headers = [ws.cell(row=1, column=c).value for c in range(1, 7)]
    assert xlsx_headers == csv_rows[0]

    # Compare data rows 1..3
    for idx, csv_row in enumerate(csv_rows[1:], start=1):
        row_idx = idx + 1
        seq_num = ws.cell(row=row_idx, column=1).value
        dt_val = ws.cell(row=row_idx, column=2).value
        proj_val = ws.cell(row=row_idx, column=3).value
        inv_val = ws.cell(row=row_idx, column=4).value
        supp_val = ws.cell(row=row_idx, column=5).value
        tax_val = ws.cell(row=row_idx, column=6).value

        assert str(seq_num) == csv_row[0]
        # Date string comparison
        if isinstance(dt_val, (datetime.date, datetime.datetime)):
            assert dt_val.strftime("%Y-%m-%d") == csv_row[1]
        else:
            assert str(dt_val) == csv_row[1]
        assert str(proj_val) == csv_row[2]
        assert str(inv_val) == csv_row[3]
        assert str(supp_val) == csv_row[4]
        assert f"{float(tax_val):.2f}" == csv_row[5]

    wb.close()


# =========================================================================
# 16. Export failure is reported correctly by the UI
# =========================================================================
def test_16_export_failure_is_reported_correctly_by_ui(qapp):
    page = ExportPage()
    rec = create_synthetic_record(1)
    page.set_data([rec], "C:/dummy/output.xlsx")

    # Simulate in-progress state
    page.btn_export.setEnabled(False)
    page.btn_export_csv.setEnabled(False)
    page.progress_frame.setVisible(True)

    page.handle_export_failure("فشل في الوصول إلى القرص")

    assert not page.progress_frame.isVisible()
    assert page.btn_export.isEnabled()
    assert page.btn_export_csv.isEnabled()


# =========================================================================
# 17. Export works with more than 183 records without any artificial row limit
# =========================================================================
def test_17_export_works_with_more_than_183_records_without_any_artificial_row_limit(tmp_path):
    exporter = XlsxExporter()
    output_path = str(tmp_path / "large_batch.xlsx")

    # Create 220 unique eligible synthetic records
    records = [
        create_synthetic_record(
            index=i,
            invoice_number=f"INV-{1000 + i}",
            tax_amount=10.0 * i
        )
        for i in range(1, 221)
    ]

    summary = exporter.export_transactions(output_path, records)
    assert summary.records_exported == 220

    wb = openpyxl.load_workbook(output_path, data_only=True)
    ws = wb["المعاملات"]
    assert ws.max_row == 221  # 1 header + 220 data rows
    assert ws.cell(row=221, column=1).value == 220
    assert ws.cell(row=221, column=4).value == "INV-1220"
    wb.close()


# =========================================================================
# 18. No existing Excel template is required
# =========================================================================
def test_18_no_existing_excel_template_is_required(tmp_path):
    # Completely empty sandbox directory with zero template files
    sandbox = tmp_path / "isolated_empty_dir"
    sandbox.mkdir()

    exporter = XlsxExporter()
    out_file = str(sandbox / "export_standalone.xlsx")

    records = [create_synthetic_record(1)]
    summary = exporter.export_transactions(out_file, records)

    assert os.path.exists(out_file)
    assert summary.records_exported == 1


# =========================================================================
# 19. No existing reference file is modified
# =========================================================================
def test_19_no_existing_reference_file_is_modified(tmp_path, synthetic_pdf_path, synthetic_template_path):
    pdf_sha_before = hashlib.sha256(open(synthetic_pdf_path, "rb").read()).hexdigest()
    xlsx_sha_before = hashlib.sha256(open(synthetic_template_path, "rb").read()).hexdigest()

    exporter = XlsxExporter()
    out_file = str(tmp_path / "safe_output.xlsx")
    records = [create_synthetic_record(1)]
    exporter.export_transactions(out_file, records)

    # Overwrite protection check: attempting to overwrite template or statement directly must fail
    with pytest.raises(XlsxProtectionError):
        exporter.export_transactions(synthetic_template_path, records, overwrite=True)

    with pytest.raises(XlsxProtectionError):
        exporter.export_transactions(synthetic_pdf_path, records, overwrite=True)

    pdf_sha_after = hashlib.sha256(open(synthetic_pdf_path, "rb").read()).hexdigest()
    xlsx_sha_after = hashlib.sha256(open(synthetic_template_path, "rb").read()).hexdigest()

    assert pdf_sha_before == pdf_sha_after
    assert xlsx_sha_before == xlsx_sha_after


# =========================================================================
# 20. Unmodified raw extraction exports ONLY cleanly eligible records (Fail-Closed)
# =========================================================================
def test_20_unmodified_raw_extraction_exports_only_eligible_records(tmp_path, synthetic_pdf_path):
    """
    Validates that feeding unmodified extracted records directly into XlsxExporter
    strictly respects fail-closed rules: only cleanly validated records are exported,
    and all unreviewed or duplicate records remain strictly blocked.
    """
    from injaz.engine import CoreExtractionEngine

    engine = CoreExtractionEngine()
    result = engine.process_statement(synthetic_pdf_path)
    assert len(result.records) == 4

    # Pass the completely untouched records directly to the exporter
    exporter = XlsxExporter()
    out_xlsx = str(tmp_path / "raw_unmodified_export.xlsx")
    summary = exporter.export_transactions(out_xlsx, result.records, overwrite=True)

    # In synthetic statement: 3 clean records exported, 1 unresolved duplicate warning blocked
    assert summary.total_records_submitted == 4
    assert summary.records_exported == 3
    assert summary.records_rejected_ineligible == 1

    wb = openpyxl.load_workbook(out_xlsx, data_only=True)
    ws = wb["المعاملات"]
    assert ws.max_row == 4  # 1 header + 3 eligible data rows
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=4, column=1).value == 3
    wb.close()


# =========================================================================
# 21. Simulated post-review scenario exports approved records with documented decisions
# =========================================================================
def test_21_simulated_post_review_scenario_exports_eligible_records(tmp_path, synthetic_pdf_path):
    """
    Validates that when a reviewer explicitly approves flagged records in the UI,
    the newly approved records become eligible and are exported, while rule-excluded
    records remain strictly blocked.
    """
    import copy
    from injaz.engine import CoreExtractionEngine

    engine = CoreExtractionEngine()
    result = engine.process_statement(synthetic_pdf_path)
    simulated_records = copy.deepcopy(result.records)

    # Add a business-rule excluded record to prove rule exclusions remain strictly blocked post-review
    rule_excl = create_synthetic_record(
        index=5,
        inclusion_eligibility=InclusionEligibility.EXCLUDED_BY_RULE,
        user_decision=UserDecision.NOT_APPLICABLE
    )
    simulated_records.append(rule_excl)

    # Documented explicit decisions:
    for r in simulated_records:
        if r.inclusion_eligibility == InclusionEligibility.EXCLUDED_BY_RULE:
            continue  # NEVER approved; remain excluded
        if any("font-spaced" in reason.lower() for reason in r.review_reasons):
            r.resolve_supplier_review(r.supplier_name, notes="Approved supplier canonical name")
        if any("single-digit" in reason.lower() for reason in r.review_reasons):
            r.resolve_field_review("Confirmed single-digit invoice")
        if r.requires_inclusion_review:
            r.user_decision = UserDecision.INCLUDED_BY_USER
        if r.duplicate_status == DuplicateStatus.POTENTIAL_DUPLICATE:
            r.resolve_duplicate_warning("Confirmed independent transaction")

    exporter = XlsxExporter()
    out_xlsx = str(tmp_path / "simulated_post_review_export.xlsx")
    summary = exporter.export_transactions(out_xlsx, simulated_records, overwrite=True)

    # 4 eligible records exported, 1 rule-excluded record rejected
    assert summary.total_records_submitted == 5
    assert summary.records_exported == 4
    assert summary.records_rejected_ineligible == 1

    wb = openpyxl.load_workbook(out_xlsx, data_only=True)
    ws = wb["المعاملات"]
    assert ws.max_row == 5  # 1 header + 4 data rows
    wb.close()

