"""
Automated Regression Test Suite for INJAZ Phase 4 Corrective Fixes:
1. Dialog contrast, light theme, and Arabic labels.
2. Excel row mapping, pre-numbered blank row handling, and sequence assignment.
3. Target worksheet verification ('تحليلي مشتريات').
4. Safe default output destination (outside references/).
5. Accurate export reporting (no false success on 0 exported).
"""
import os
import sys
import hashlib
import openpyxl
import pytest

os.environ["QT_QPA_PLATFORM"] = "offscreen"

from PySide6.QtWidgets import QApplication, QDialog, QMessageBox
from PySide6.QtCore import Qt

from injaz.models import (
    TransactionRecord,
    ExtractionStatus,
    InclusionEligibility,
    UserDecision,
    DuplicateStatus,
)
from injaz.excel_writer import ExcelPopulationEngine, ExportSummary, ExcelPopulationError
from injaz.ui.components.dialogs import (
    ExportConfirmDialog, ExportSuccessDialog, ExportSkippedDialog
)
from injaz.ui.pages.new_process_page import NewProcessPage
from injaz.ui.pages.export_page import ExportPage


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance()
    if app is None:
        app = QApplication(sys.argv)
    app.setLayoutDirection(Qt.RightToLeft)
    return app


@pytest.fixture(autouse=True)
def mock_dialogs(monkeypatch):
    """Mocks blocking dialog execution for headless testing."""
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QDialog, "exec", lambda self: QDialog.Accepted)
    monkeypatch.setattr(QDialog, "exec_", lambda self: QDialog.Accepted)


def make_test_record(index: int, invoice: str = "500", tax: float = 75.0) -> TransactionRecord:
    return TransactionRecord(
        page_number=1,
        record_index=index,
        date="2026-05-10",
        transaction_type="قيد يدوي",
        raw_description=f"فاتورة {invoice} شركة الاختبار",
        reference=f"REF-{index}",
        debit=tax,
        credit=0.0,
        balance="1000.0",
        project="مشروع تجريبي",
        invoice_number=invoice,
        supplier_name="شركة الاختبار والتوريد",
        tax_amount=tax,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE,
    )


def create_blank_prenumbered_template(filepath: str, num_placeholder_rows: int = 10):
    """Creates a synthetic workbook mimicking the Injaz purchases template with pre-numbered blank rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "تحليلي مشتريات"

    # Row 1: Title
    ws.cell(1, 1).value = "تحليلي مشتريات 2026"

    # Row 2: Headers
    headers = ["م", "التاريخ", "المشروع", "رقم الفاتورة", "المورد", "القيمة", "الضريبة", "الإجمالي"]
    for c, h in enumerate(headers, start=1):
        ws.cell(2, c).value = h

    # Pre-numbered blank rows (rows 3 to 2 + num_placeholder_rows)
    for i in range(1, num_placeholder_rows + 1):
        r = 2 + i
        ws.cell(r, 1).value = i  # Sequence number only
        ws.cell(r, 6).value = f"=G{r}/0.15"  # Value formula
        ws.cell(r, 8).value = f"=F{r}+G{r}"  # Total formula

    # Summary row
    sum_row = 3 + num_placeholder_rows
    ws.cell(sum_row, 1).value = "الإجمالي"
    ws.cell(sum_row, 6).value = f"=SUM(F3:F{sum_row-1})"
    ws.cell(sum_row, 7).value = f"=SUM(G3:G{sum_row-1})"
    ws.cell(sum_row, 8).value = f"=SUM(H3:H{sum_row-1})"

    # Other sheets
    wb.create_sheet("Sheet1")
    wb.create_sheet("ملخص")

    wb.save(filepath)
    wb.close()


# ---------------------------------------------------------------------------
# Test 1: Confirmation dialog light theme and Arabic labels
# ---------------------------------------------------------------------------
def test_1_confirmation_dialog_light_theme_and_arabic_labels(qapp):
    dlg = ExportConfirmDialog(
        eligible_count=15,
        output_path="C:/Exports/out.xlsx",
        template_path="C:/Templates/tpl.xlsx"
    )
    assert dlg.layoutDirection() == Qt.RightToLeft
    assert dlg.windowTitle() == "تأكيد عملية التصدير"
    assert dlg.btn_confirm.text() == "تأكيد التصدير"
    assert dlg.btn_cancel.text() == "إلغاء"
    # Verify stylesheet specifies light surface and Deep Navy / Teal palette
    ss = dlg.styleSheet()
    assert "#FFFFFF" in ss
    assert "#16324F" in ss
    assert "#1D9A8A" in ss
    dlg.close()


# ---------------------------------------------------------------------------
# Test 2: UI calls writer with actual eligible records after confirmation
# ---------------------------------------------------------------------------
def test_2_ui_calls_writer_with_actual_eligible_records(qapp):
    page = ExportPage()
    rec1 = make_test_record(1)
    rec2 = make_test_record(2)
    page.set_data([rec1, rec2], "tpl.xlsx", "out.xlsx")

    emitted = []
    page.export_started.connect(lambda t, o, r: emitted.append((t, o, r)))
    page._on_export_clicked()

    assert len(emitted) == 1
    t_path, o_path, recs = emitted[0]
    assert t_path == "tpl.xlsx"
    assert o_path == "out.xlsx"
    assert len(recs) == 2


# ---------------------------------------------------------------------------
# Test 3: Writer exception reported in UI
# ---------------------------------------------------------------------------
def test_3_writer_exception_reported_in_ui(qapp, monkeypatch):
    page = ExportPage()
    errors_shown = []
    monkeypatch.setattr(QMessageBox, "critical", lambda parent, title, msg: errors_shown.append(msg))

    page.progress_frame.setVisible(True)
    page.btn_export.setEnabled(False)

    page.handle_export_failure("خطأ في بنية القالب المحدد")
    assert not page.progress_frame.isVisible()
    assert page.btn_export.isEnabled()
    assert len(errors_shown) == 1
    assert "خطأ في بنية القالب" in errors_shown[0]


# ---------------------------------------------------------------------------
# Test 4: Successful export creates real output workbook containing records
# ---------------------------------------------------------------------------
def test_4_successful_export_creates_real_output_workbook(tmp_path):
    tpl_file = tmp_path / "test_template.xlsx"
    create_blank_prenumbered_template(str(tpl_file), num_placeholder_rows=5)
    out_file = tmp_path / "output_test.xlsx"

    engine = ExcelPopulationEngine()
    recs = [make_test_record(1, "101", 50.0), make_test_record(2, "102", 75.0)]

    summary = engine.export_transactions(
        template_path=str(tpl_file),
        output_path=str(out_file),
        records=recs,
        overwrite=True
    )

    assert summary.records_exported == 2
    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 0

    # Reopen and verify data
    wb = openpyxl.load_workbook(str(out_file), data_only=False)
    ws = wb["تحليلي مشتريات"]
    # Row 3 (first transaction row)
    assert ws.cell(3, 1).value == 1  # Sequence
    assert ws.cell(3, 3).value == "مشروع تجريبي"  # Project
    assert ws.cell(3, 4).value == 101  # Invoice
    assert ws.cell(3, 5).value == "شركة الاختبار والتوريد"  # Supplier
    assert ws.cell(3, 7).value == 50.0  # Tax

    # Row 4 (second transaction row)
    assert ws.cell(4, 1).value == 2
    assert ws.cell(4, 4).value == 102
    assert ws.cell(4, 7).value == 75.0
    wb.close()


# ---------------------------------------------------------------------------
# Test 5: Target worksheet is explicitly 'تحليلي مشتريات'
# ---------------------------------------------------------------------------
def test_5_target_worksheet_is_explicitly_purchases_sheet(tmp_path):
    tpl_file = tmp_path / "template.xlsx"
    create_blank_prenumbered_template(str(tpl_file))
    out_file = tmp_path / "out.xlsx"

    engine = ExcelPopulationEngine()
    engine.export_transactions(
        template_path=str(tpl_file),
        output_path=str(out_file),
        records=[make_test_record(1)],
        overwrite=True
    )

    wb = openpyxl.load_workbook(str(out_file))
    assert "تحليلي مشتريات" in wb.sheetnames
    ws_other = wb["Sheet1"]
    assert ws_other.max_row == 1 or ws_other.cell(1, 1).value is None
    wb.close()


# ---------------------------------------------------------------------------
# Test 6: Correct column field mapping
# ---------------------------------------------------------------------------
def test_6_correct_column_field_mapping(tmp_path):
    tpl_file = tmp_path / "template.xlsx"
    create_blank_prenumbered_template(str(tpl_file))
    out_file = tmp_path / "out.xlsx"

    engine = ExcelPopulationEngine()
    engine.export_transactions(
        template_path=str(tpl_file),
        output_path=str(out_file),
        records=[make_test_record(1, invoice="9988", tax=120.0)],
        overwrite=True
    )

    wb = openpyxl.load_workbook(str(out_file))
    ws = wb["تحليلي مشتريات"]
    # Column A: Sequence
    assert ws.cell(3, 1).value == 1
    # Column B: Date
    assert ws.cell(3, 2).value is not None
    # Column C: Project
    assert ws.cell(3, 3).value == "مشروع تجريبي"
    # Column D: Invoice
    assert ws.cell(3, 4).value == 9988
    # Column E: Supplier
    assert ws.cell(3, 5).value == "شركة الاختبار والتوريد"
    # Column G: Tax
    assert ws.cell(3, 7).value == 120.0
    wb.close()


# ---------------------------------------------------------------------------
# Test 7: Pre-numbered rows with no data are treated as empty slots
# ---------------------------------------------------------------------------
def test_7_prenumbered_rows_treated_as_empty_slots(tmp_path):
    tpl_file = tmp_path / "template.xlsx"
    # Create template with 8 pre-numbered blank slots
    create_blank_prenumbered_template(str(tpl_file), num_placeholder_rows=8)

    engine = ExcelPopulationEngine()
    wb = openpyxl.load_workbook(str(tpl_file))
    ws = wb["تحليلي مشتريات"]
    header_idx, col_map = engine.locate_headers(ws)
    last_data_row, summary_row, existing_keys, max_seq = engine.inspect_existing_state(
        ws, col_map, header_idx
    )
    wb.close()

    # Crucial assertion: last_data_row must be header row (2), NOT row 10!
    assert last_data_row == 2
    assert max_seq == 0
    assert len(existing_keys) == 0


# ---------------------------------------------------------------------------
# Test 8: Numbering starts correctly for blank and continues for populated
# ---------------------------------------------------------------------------
def test_8_numbering_starts_correctly_and_continues(tmp_path):
    tpl_file = tmp_path / "template.xlsx"
    create_blank_prenumbered_template(str(tpl_file), num_placeholder_rows=5)
    out_file1 = tmp_path / "out1.xlsx"

    engine = ExcelPopulationEngine()
    # Batch 1: 2 records into blank template -> sequence 1, 2
    engine.export_transactions(
        template_path=str(tpl_file),
        output_path=str(out_file1),
        records=[make_test_record(1, invoice="101"), make_test_record(2, invoice="102")],
        overwrite=True
    )

    wb1 = openpyxl.load_workbook(str(out_file1))
    ws1 = wb1["تحليلي مشتريات"]
    assert ws1.cell(3, 1).value == 1
    assert ws1.cell(4, 1).value == 2
    wb1.close()

    # Batch 2: Export 1 more record using out1 as template -> sequence continues to 3
    out_file2 = tmp_path / "out2.xlsx"
    rec3 = make_test_record(3, invoice="777", tax=33.0)
    engine.export_transactions(
        template_path=str(out_file1),
        output_path=str(out_file2),
        records=[rec3],
        overwrite=True
    )

    wb2 = openpyxl.load_workbook(str(out_file2))
    ws2 = wb2["تحليلي مشتريات"]
    assert ws2.cell(3, 1).value == 1
    assert ws2.cell(4, 1).value == 2
    assert ws2.cell(5, 1).value == 3
    wb2.close()


# ---------------------------------------------------------------------------
# Test 9: Formulas in Value and Total remain intact
# ---------------------------------------------------------------------------
def test_9_formulas_in_value_and_total_remain_intact(tmp_path):
    tpl_file = tmp_path / "template.xlsx"
    create_blank_prenumbered_template(str(tpl_file))
    out_file = tmp_path / "out.xlsx"

    engine = ExcelPopulationEngine()
    engine.export_transactions(
        template_path=str(tpl_file),
        output_path=str(out_file),
        records=[make_test_record(1)],
        overwrite=True
    )

    wb = openpyxl.load_workbook(str(out_file), data_only=False)
    ws = wb["تحليلي مشتريات"]
    c_val = ws.cell(3, 6).value
    c_tot = ws.cell(3, 8).value
    assert isinstance(c_val, str) and c_val.startswith("=")
    assert isinstance(c_tot, str) and c_tot.startswith("=")
    wb.close()


# ---------------------------------------------------------------------------
# Test 10: Source template SHA-256 remains completely unchanged
# ---------------------------------------------------------------------------
def test_10_source_template_sha256_remains_unchanged(tmp_path):
    tpl_file = tmp_path / "template.xlsx"
    create_blank_prenumbered_template(str(tpl_file))

    def sha(p):
        with open(p, "rb") as f:
            return hashlib.sha256(f.read()).hexdigest()

    h_before = sha(str(tpl_file))
    out_file = tmp_path / "out.xlsx"

    engine = ExcelPopulationEngine()
    engine.export_transactions(
        template_path=str(tpl_file),
        output_path=str(out_file),
        records=[make_test_record(1)],
        overwrite=True
    )

    h_after = sha(str(tpl_file))
    assert h_before == h_after


# ---------------------------------------------------------------------------
# Test 11: Default output destination is outside references/ directory
# ---------------------------------------------------------------------------
def test_11_default_output_destination_outside_references(qapp, tmp_path):
    page = NewProcessPage()
    # Select pdf in a references-like dir
    pdf_in_refs = tmp_path / "references" / "sample.pdf"
    pdf_in_refs.parent.mkdir(parents=True, exist_ok=True)
    pdf_in_refs.write_text("pdf")

    page._on_pdf_selected(str(pdf_in_refs))
    # Must NOT be inside references/
    assert "references" not in page.output_path.lower()
    assert "INJAZ" in page.output_path or "Documents" in page.output_path


# ---------------------------------------------------------------------------
# Test 12: Failed or skipped export never displays false success message
# ---------------------------------------------------------------------------
def test_12_failed_or_skipped_export_never_displays_false_success(qapp, monkeypatch):
    page = ExportPage()

    success_opened = []
    skipped_opened = []

    monkeypatch.setattr(ExportSuccessDialog, "exec", lambda self: success_opened.append(True))
    monkeypatch.setattr(ExportSkippedDialog, "exec", lambda self: skipped_opened.append(True))

    # Case A: 0 exported because all 5 were already present
    summary_all_skipped = ExportSummary(
        total_records_submitted=5,
        records_exported=0,
        records_skipped_already_present=5,
        records_rejected_ineligible=0,
        output_path="C:/out.xlsx"
    )
    page.handle_export_success(summary_all_skipped)
    assert len(success_opened) == 0
    assert len(skipped_opened) == 1

    # Case B: 3 records actually exported
    summary_exported = ExportSummary(
        total_records_submitted=3,
        records_exported=3,
        records_skipped_already_present=0,
        records_rejected_ineligible=0,
        output_path="C:/out.xlsx"
    )
    page.handle_export_success(summary_exported)
    assert len(success_opened) == 1
