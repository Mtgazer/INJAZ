"""
Comprehensive Automated Test Suite for Injaz (إنجاز) — Phase 3.
Validates Excel Population Engine, template protection, fail-closed eligibility,
formula/style preservation, duplicate prevention, and workbook integrity.
"""
import os
import shutil
import hashlib
import tempfile
import datetime
from typing import List
import pytest
import openpyxl

from injaz.excel_writer import (
    ExcelPopulationEngine,
    ColumnMapping,
    TemplateProtectionError,
    IncompatibleTemplateError,
    IneligibleExportError,
    WorkbookValidationError,
)
from injaz.models import (
    TransactionRecord,
    ExtractionStatus,
    InclusionEligibility,
    UserDecision,
    DuplicateStatus,
)


def get_file_sha256(filepath: str) -> str:
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


@pytest.fixture
def reference_template_path():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    path = os.path.join(base_dir, "tests", "fixtures", "synthetic_template.xlsx")
    assert os.path.exists(path), f"Synthetic template not found at {path}"
    return path


@pytest.fixture
def approved_sample_record():
    return TransactionRecord(
        page_number=1,
        record_index=101,
        date="2026-07-01",
        transaction_type="قيد يدوي",
        raw_description="شراء مواد بناء لمشروع صيل 14",
        reference="REF-7788",
        debit=150.0,
        credit=0.0,
        balance="150.0",
        project="صيل 14",
        invoice_number="5544",
        supplier_name="شركة الرواف للتجارة والمقاولات",
        tax_amount=22.5,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE,
    )


# =============================================================================
# 1. Correct mapping when input column order changes
# =============================================================================

def test_1_reordered_columns_mapping(tmp_path):
    """Verifies that columns are located dynamically by header name, even when column order is shuffled."""
    custom_template = str(tmp_path / "shuffled_template.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "تحليلي مشتريات"

    # Row 1: Title
    ws.append(["كشف مشتريات اختباري"])
    # Row 2: Shuffled Headers
    ws.append(["الضريبة", "المورد", "م", "رقم الفاتورة", "المشروع", "التاريخ", "القيمة", "الإجمالي"])
    # Row 3: Initial Data
    ws.append([15.0, "المورد الأول", 1, 1001, "مشروع أ", datetime.date(2026, 6, 1), "=A3/0.15", "=G3+A3"])
    wb.save(custom_template)
    wb.close()

    out_file = str(tmp_path / "shuffled_output.xlsx")
    rec = TransactionRecord(
        page_number=1, record_index=1, date="2026-07-02",
        transaction_type="قيد يدوي", raw_description="مواد", reference="R1",
        debit=30.0, credit=0.0, balance="30.0", project="مشروع ب",
        invoice_number="1002", supplier_name="المورد الثاني", tax_amount=30.0,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE
    )

    engine = ExcelPopulationEngine(sheet_name="تحليلي مشتريات", header_row=2)
    summary = engine.export_transactions(custom_template, out_file, [rec])

    assert summary.records_exported == 1
    assert summary.rows_appended_indices == [4]

    wb_res = openpyxl.load_workbook(out_file, data_only=False)
    ws_res = wb_res["تحليلي مشتريات"]
    # Check that data went to the shuffled column positions:
    # Col 1: الضريبة
    assert abs(float(ws_res.cell(4, 1).value) - 30.0) < 0.001
    # Col 2: المورد
    assert ws_res.cell(4, 2).value == "المورد الثاني"
    # Col 3: م (Seq)
    assert ws_res.cell(4, 3).value == 2
    # Col 4: رقم الفاتورة
    assert ws_res.cell(4, 4).value == 1002
    # Col 5: المشروع
    assert ws_res.cell(4, 5).value == "مشروع ب"
    # Col 6: التاريخ
    val_dt = ws_res.cell(4, 6).value
    assert str(val_dt)[:10] == "2026-07-02"
    # Col 7 & 8: Formulas
    assert str(ws_res.cell(4, 7).value).startswith("=")
    assert str(ws_res.cell(4, 8).value).startswith("=")
    wb_res.close()


# =============================================================================
# 2. Missing worksheet or required header
# =============================================================================

def test_2_missing_worksheet_and_required_header(tmp_path):
    """Verifies that missing worksheets or required headers raise IncompatibleTemplateError."""
    # A. Missing worksheet
    bad_sheet_template = str(tmp_path / "bad_sheet.xlsx")
    wb1 = openpyxl.Workbook()
    wb1.active.title = "Sheet_Unknown"
    wb1.save(bad_sheet_template)
    wb1.close()

    engine = ExcelPopulationEngine(sheet_name="تحليلي مشتريات")
    with pytest.raises(IncompatibleTemplateError, match="Worksheet 'تحليلي مشتريات' not found"):
        engine.export_transactions(bad_sheet_template, str(tmp_path / "out1.xlsx"), [])

    # B. Missing required header (e.g. missing 'الضريبة')
    bad_header_template = str(tmp_path / "bad_header.xlsx")
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "تحليلي مشتريات"
    ws2.append(["العنوان"])
    # Missing 'الضريبة'
    ws2.append(["م", "التاريخ", "المشروع", "رقم الفاتورة", "المورد"])
    wb2.save(bad_header_template)
    wb2.close()

    with pytest.raises(IncompatibleTemplateError, match="missing required header"):
        engine.export_transactions(bad_header_template, str(tmp_path / "out2.xlsx"), [])


# =============================================================================
# 3. Successful population of approved records
# =============================================================================

def test_3_successful_population_approved_records(reference_template_path, tmp_path, approved_sample_record):
    """Verifies successful population of approved records into template."""
    out_file = str(tmp_path / "test3_out.xlsx")
    engine = ExcelPopulationEngine()

    summary = engine.export_transactions(
        template_path=reference_template_path,
        output_path=out_file,
        records=[approved_sample_record]
    )

    assert summary.records_exported == 1
    assert summary.records_skipped_already_present == 0
    assert summary.records_rejected_ineligible == 0
    assert summary.rows_appended_indices == [192]

    # Verify content in row 192
    wb = openpyxl.load_workbook(out_file, data_only=False)
    ws = wb["تحليلي مشتريات"]
    assert str(ws.cell(192, 2).value)[:10] == "2026-07-01"
    assert ws.cell(192, 3).value == "صيل 14"
    assert ws.cell(192, 4).value == 5544
    assert ws.cell(192, 5).value == "شركة الرواف للتجارة والمقاولات"
    assert abs(float(ws.cell(192, 7).value) - 22.5) < 0.001
    wb.close()


# =============================================================================
# 4. Rejection of excluded records
# =============================================================================

def test_4_rejection_of_excluded_records(reference_template_path, tmp_path):
    """Verifies that records classified as EXCLUDED_BY_RULE are never exported."""
    out_file = str(tmp_path / "test4_out.xlsx")
    rec = TransactionRecord(
        page_number=1, record_index=201, date="2026-04-10",
        transaction_type="قيد يدوي", raw_description="سداد فاتورة كهرباء", reference="REF-ELEC",
        debit=50.0, credit=0.0, balance="50.0", project="صيل 13",
        invoice_number="9988", supplier_name="شركة الكهرباء السعودية", tax_amount=7.5,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.EXCLUDED_BY_RULE,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE
    )

    engine = ExcelPopulationEngine()
    # Default behavior: tracks rejection in summary
    summary = engine.export_transactions(reference_template_path, out_file, [rec])
    assert summary.records_exported == 0
    assert summary.records_rejected_ineligible == 1
    assert "EXCLUDED_BY_RULE" in summary.rejection_details[0]["reason"]

    # With fail_on_ineligible=True: raises IneligibleExportError
    with pytest.raises(IneligibleExportError, match="EXCLUDED_BY_RULE"):
        engine.export_transactions(
            reference_template_path, out_file, [rec], overwrite=True, fail_on_ineligible=True
        )


# =============================================================================
# 5. Rejection of unresolved field review
# =============================================================================

def test_5_rejection_of_unresolved_field_review(reference_template_path, tmp_path):
    """Verifies that records with unresolved manual review cannot be exported until resolved."""
    out_file = str(tmp_path / "test5_out.xlsx")
    rec = TransactionRecord(
        page_number=1, record_index=301, date="2026-05-15",
        transaction_type="قيد يدوي", raw_description="توريد مواد", reference="REF-301",
        debit=200.0, credit=0.0, balance="200.0", project="صيل 15",
        invoice_number="3322", supplier_name="مؤس سة ال ر واف", tax_amount=30.0,
        extraction_status=ExtractionStatus.MANUAL_REVIEW,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE,
        review_reasons=["UNCERTAIN_SUPPLIER_FONT_SPACING"]
    )

    engine = ExcelPopulationEngine()
    summary1 = engine.export_transactions(reference_template_path, out_file, [rec])
    assert summary1.records_exported == 0
    assert summary1.records_rejected_ineligible == 1

    # Now human resolves field review
    rec.resolve_field_review("Verified visually against physical voucher")
    assert rec.is_export_eligible is True

    summary2 = engine.export_transactions(reference_template_path, out_file, [rec], overwrite=True)
    assert summary2.records_exported == 1
    assert summary2.records_rejected_ineligible == 0


# =============================================================================
# 6. Rejection of missing inclusion decisions
# =============================================================================

def test_6_rejection_of_missing_inclusion_decisions(reference_template_path, tmp_path):
    """Verifies that records requiring review cannot be exported without explicit inclusion approval."""
    out_file = str(tmp_path / "test6_out.xlsx")
    rec = TransactionRecord(
        page_number=1, record_index=401, date="2026-05-20",
        transaction_type="قيد يدوي", raw_description="معاملة مركز رئيسي استشارية", reference="REF-HO",
        debit=500.0, credit=0.0, balance="500.0", project="المركز الرئيسي",
        invoice_number="7711", supplier_name="مكتب الاستشارات الهندسية", tax_amount=75.0,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.REQUIRES_USER_REVIEW,
        user_decision=UserDecision.PENDING,
        duplicate_status=DuplicateStatus.UNIQUE
    )

    engine = ExcelPopulationEngine()
    # While PENDING: rejected
    summary1 = engine.export_transactions(reference_template_path, out_file, [rec])
    assert summary1.records_exported == 0
    assert summary1.records_rejected_ineligible == 1

    # User explicitly excludes: rejected
    rec.user_decision = UserDecision.EXCLUDED_BY_USER
    summary2 = engine.export_transactions(reference_template_path, out_file, [rec], overwrite=True)
    assert summary2.records_exported == 0

    # User explicitly includes: approved for export
    rec.user_decision = UserDecision.INCLUDED_BY_USER
    summary3 = engine.export_transactions(reference_template_path, out_file, [rec], overwrite=True)
    assert summary3.records_exported == 1


# =============================================================================
# 7. Rejection of unresolved duplicate warnings
# =============================================================================

def test_7_rejection_of_unresolved_duplicate_warnings(reference_template_path, tmp_path):
    """Verifies that potential duplicates are held until explicitly resolved by user."""
    out_file = str(tmp_path / "test7_out.xlsx")
    rec = TransactionRecord(
        page_number=1, record_index=501, date="2026-06-01",
        transaction_type="قيد يدوي", raw_description="شراء حديد تسليح", reference="REF-DUP",
        debit=1000.0, credit=0.0, balance="1000.0", project="صيل 16",
        invoice_number="8899", supplier_name="شركة سابك للمعادن", tax_amount=150.0,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.POTENTIAL_DUPLICATE,
        duplicate_key="2026-06-01|صيل16|8899|شركةسابكللمعادن|150.00"
    )

    engine = ExcelPopulationEngine()
    # Unresolved: rejected
    summary1 = engine.export_transactions(reference_template_path, out_file, [rec])
    assert summary1.records_exported == 0
    assert summary1.records_rejected_ineligible == 1

    # Human auditor resolves duplicate warning
    rec.resolve_duplicate_warning("Verified as second delivery tranche")
    assert rec.is_export_eligible is True

    summary2 = engine.export_transactions(reference_template_path, out_file, [rec], overwrite=True)
    assert summary2.records_exported == 1


# =============================================================================
# 8. Preservation of all existing formula cells, especially Value and Total
# =============================================================================

def test_8_preservation_of_existing_formula_cells(reference_template_path, tmp_path, approved_sample_record):
    """Verifies that Value and Total formulas in existing and new rows are strictly preserved."""
    out_file = str(tmp_path / "test8_out.xlsx")
    engine = ExcelPopulationEngine()

    summary = engine.export_transactions(reference_template_path, out_file, [approved_sample_record])
    assert summary.records_exported == 1

    wb = openpyxl.load_workbook(out_file, data_only=False)
    ws = wb["تحليلي مشتريات"]

    # Verify historical rows (e.g. Row 3 and Row 191) have formulas intact
    val_f3 = ws.cell(3, 6).value
    tot_h3 = ws.cell(3, 8).value
    assert isinstance(val_f3, str) and val_f3.startswith("=")
    assert isinstance(tot_h3, str) and tot_h3.startswith("=")

    val_f191 = ws.cell(191, 6).value
    tot_h191 = ws.cell(191, 8).value
    assert isinstance(val_f191, str) and val_f191.startswith("=")
    assert isinstance(tot_h191, str) and tot_h191.startswith("=")

    # Verify populated row 192 has formula preserved
    val_f192 = ws.cell(192, 6).value
    tot_h192 = ws.cell(192, 8).value
    assert isinstance(val_f192, str) and val_f192.startswith("=")
    assert isinstance(tot_h192, str) and tot_h192.startswith("=")

    # Verify summary row (row 375) has SUBTOTAL formulas intact
    val_subtotal_val = ws.cell(375, 6).value
    val_subtotal_tax = ws.cell(375, 7).value
    val_subtotal_tot = ws.cell(375, 8).value
    assert "=SUBTOTAL" in str(val_subtotal_val)
    assert "=SUBTOTAL" in str(val_subtotal_tax)
    assert "=SUBTOTAL" in str(val_subtotal_tot)

    wb.close()


# =============================================================================
# 9. Preservation of other worksheets and existing data
# =============================================================================

def test_9_preservation_of_other_worksheets_and_existing_data(reference_template_path, tmp_path, approved_sample_record):
    """Verifies that other worksheets ('Sheet1', 'ملخص') and existing rows remain untouched."""
    out_file = str(tmp_path / "test9_out.xlsx")
    engine = ExcelPopulationEngine()

    engine.export_transactions(reference_template_path, out_file, [approved_sample_record])

    wb = openpyxl.load_workbook(out_file, data_only=False)
    assert "Sheet1" in wb.sheetnames
    assert "ملخص" in wb.sheetnames

    # Check Sheet1 data
    ws_sheet1 = wb["Sheet1"]
    assert ws_sheet1.cell(4, 2).value == "شركة النماذج الافتراضية"
    assert ws_sheet1.cell(4, 3).value == "3999999999"

    # Check ملخص data
    ws_summary = wb["ملخص"]
    assert ws_summary.cell(1, 1).value == "المشروع"
    assert ws_summary.cell(1, 2).value == "إجمالي الضريبة"

    wb.close()


# =============================================================================
# 10. Preservation of relevant formatting and number formats
# =============================================================================

def test_10_preservation_of_formatting_and_number_formats(reference_template_path, tmp_path, approved_sample_record):
    """Verifies that styling, font names, borders, and number formats match template."""
    out_file = str(tmp_path / "test10_out.xlsx")
    engine = ExcelPopulationEngine()

    engine.export_transactions(reference_template_path, out_file, [approved_sample_record])

    wb = openpyxl.load_workbook(out_file, data_only=False)
    ws = wb["تحليلي مشتريات"]
    c192_tax = ws.cell(192, 7)
    c192_dt = ws.cell(192, 2)
    c192_proj = ws.cell(192, 3)

    assert "Simplified Arabic" in (c192_proj.font.name or "")
    assert c192_dt.number_format == "yyyy-mm-dd"
    assert "0.00" in c192_tax.number_format
    wb.close()


# =============================================================================
# 11. Correct behavior when output file already exists
# =============================================================================

def test_11_output_file_already_exists_behavior(reference_template_path, tmp_path, approved_sample_record):
    """Verifies safe refusal when overwrite=False and backup creation when overwrite=True."""
    out_file = str(tmp_path / "test11_existing.xlsx")
    # Pre-create output file
    shutil.copy2(reference_template_path, out_file)

    engine = ExcelPopulationEngine()
    # A. Safe refusal when overwrite=False
    with pytest.raises(FileExistsError, match="already exists"):
        engine.export_transactions(reference_template_path, out_file, [approved_sample_record], overwrite=False)

    # B. Explicit overwrite with backup
    summary = engine.export_transactions(
        reference_template_path, out_file, [approved_sample_record], overwrite=True, backup_existing=True
    )
    assert summary.records_exported == 1
    assert os.path.exists(out_file)
    assert os.path.exists(out_file + ".bak")


# =============================================================================
# 12. Rejection when source and destination paths are identical
# =============================================================================

def test_12_source_and_destination_identical_rejection(reference_template_path):
    """Verifies strict refusal when output path points to source template."""
    engine = ExcelPopulationEngine()
    with pytest.raises(TemplateProtectionError, match="Output path resolves to the source template"):
        engine.export_transactions(
            template_path=reference_template_path,
            output_path=reference_template_path,
            records=[]
        )


# =============================================================================
# 13. Validation failure during temporary output generation
# =============================================================================

def test_13_validation_failure_during_temporary_output(reference_template_path, tmp_path, approved_sample_record, monkeypatch):
    """Verifies that validation failure cleans up temporary files and leaves output untouched."""
    out_file = str(tmp_path / "test13_failed.xlsx")
    engine = ExcelPopulationEngine()

    def mock_validate(*args, **kwargs):
        raise WorkbookValidationError("Simulated corrupted formula check")

    monkeypatch.setattr(engine, "_validate_generated_workbook", mock_validate)

    with pytest.raises(WorkbookValidationError, match="Simulated corrupted formula check"):
        engine.export_transactions(reference_template_path, out_file, [approved_sample_record])

    # Assert destination file was NOT finalized
    assert not os.path.exists(out_file)

    # Assert no stray temporary files remained in the directory
    temp_files = [f for f in os.listdir(tmp_path) if f.startswith("injaz_export_")]
    assert len(temp_files) == 0


# =============================================================================
# 14. Source template SHA-256 remains unchanged
# =============================================================================

def test_14_source_template_sha256_immutability(reference_template_path, tmp_path, approved_sample_record):
    """Verifies that the reference template file is never modified by any engine operation."""
    initial_hash = get_file_sha256(reference_template_path)

    # Run population operations
    out_file = str(tmp_path / "test14_out.xlsx")
    engine = ExcelPopulationEngine()
    engine.export_transactions(reference_template_path, out_file, [approved_sample_record])

    final_hash = get_file_sha256(reference_template_path)
    assert initial_hash == final_hash, "CRITICAL ERROR: Source template was altered!"


# =============================================================================
# 15. Empty approved-record batches
# =============================================================================

def test_15_empty_approved_record_batch(reference_template_path, tmp_path):
    """Verifies that empty record batches are handled safely without modifying template data."""
    out_file = str(tmp_path / "test15_empty.xlsx")
    engine = ExcelPopulationEngine()

    summary = engine.export_transactions(reference_template_path, out_file, records=[])
    assert summary.total_records_submitted == 0
    assert summary.records_exported == 0
    assert os.path.exists(out_file)

    # Verify output workbook is intact and valid
    wb = openpyxl.load_workbook(out_file, data_only=False)
    assert wb["تحليلي مشتريات"].max_row >= 191
    wb.close()


# =============================================================================
# 16. Prevention of duplicate appends across repeated processing runs
# =============================================================================

def test_16_prevention_of_duplicate_appends_across_repeated_runs(reference_template_path, tmp_path):
    """Verifies that records already present in the target workbook are skipped on subsequent runs."""
    out_file = str(tmp_path / "test16_run1.xlsx")
    engine = ExcelPopulationEngine()

    r1 = TransactionRecord(
        page_number=1, record_index=601, date="2026-07-10",
        transaction_type="قيد يدوي", raw_description="شراء كابلات", reference="REF-CABLE",
        debit=300.0, credit=0.0, balance="300.0", project="صيل 13",
        invoice_number="7001", supplier_name="شركة الكابلات المتحدة", tax_amount=45.0,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE
    )
    r2 = TransactionRecord(
        page_number=1, record_index=602, date="2026-07-11",
        transaction_type="قيد يدوي", raw_description="شراء أنابيب", reference="REF-PIPE",
        debit=400.0, credit=0.0, balance="400.0", project="صيل 14",
        invoice_number="7002", supplier_name="شركة الأنابيب الحديثة", tax_amount=60.0,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE
    )

    # Run 1: Export [r1, r2]
    summary1 = engine.export_transactions(reference_template_path, out_file, [r1, r2])
    assert summary1.records_exported == 2
    assert summary1.records_skipped_already_present == 0
    assert summary1.rows_appended_indices == [192, 193]

    # Run 2: Submitting [r1, r2, r3] to an update using Run 1's file as source template
    out_file_run2 = str(tmp_path / "test16_run2.xlsx")
    r3 = TransactionRecord(
        page_number=1, record_index=603, date="2026-07-12",
        transaction_type="قيد يدوي", raw_description="شراء أسمنت", reference="REF-CEMENT",
        debit=500.0, credit=0.0, balance="500.0", project="صيل 15",
        invoice_number="7003", supplier_name="مصنع الأسمنت الوطني", tax_amount=75.0,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE
    )

    summary2 = engine.export_transactions(out_file, out_file_run2, [r1, r2, r3])
    # r1 and r2 are detected as already existing in the workbook, only r3 is appended!
    assert summary2.total_records_submitted == 3
    assert summary2.records_exported == 1
    assert summary2.records_skipped_already_present == 2
    assert summary2.rows_appended_indices == [194]

    # Verify that row 194 is r3
    wb = openpyxl.load_workbook(out_file_run2, data_only=False)
    ws = wb["تحليلي مشتريات"]
    assert ws.cell(194, 4).value == 7003
    assert ws.cell(194, 5).value == "مصنع الأسمنت الوطني"
    wb.close()


# =============================================================================
# 17. 184-Record Overflow Boundary Test
# =============================================================================

def test_17_overflow_184_records_boundary_test(reference_template_path, tmp_path):
    """
    Boundary test: exports 184 unique approved synthetic transactions to fill all
    183 placeholder rows (rows 192..374) and insert 1 overflow row at row 375.
    Verifies Table6.ref expansion to A2:L375, shifting of Grand-Total to row 376,
    preservation of summary formulas, formatting, and successful reopening.
    """
    out_file = str(tmp_path / "test17_overflow_184.xlsx")
    engine = ExcelPopulationEngine()

    # Generate 184 unique approved synthetic records
    records: List[TransactionRecord] = []
    base_date = datetime.date(2026, 7, 1)
    for i in range(1, 185):
        dt = (base_date + datetime.timedelta(days=i)).isoformat()
        records.append(TransactionRecord(
            page_number=1,
            record_index=1000 + i,
            date=dt,
            transaction_type="قيد يدوي",
            raw_description=f"شراء مواد دفعة {i}",
            reference=f"REF-{1000 + i}",
            debit=100.0 + i,
            credit=0.0,
            balance="100.0",
            project=f"مشروع {i % 5 + 1}",
            invoice_number=f"{80000 + i}",
            supplier_name=f"مورد تجريبي {i}",
            tax_amount=15.0 + (i * 0.1),
            extraction_status=ExtractionStatus.VALIDATED,
            inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
            user_decision=UserDecision.NOT_APPLICABLE,
            duplicate_status=DuplicateStatus.UNIQUE
        ))

    summary = engine.export_transactions(
        template_path=reference_template_path,
        output_path=out_file,
        records=records,
        overwrite=False
    )

    # 1. Verify export counts
    assert summary.total_records_submitted == 184
    assert summary.records_exported == 184
    assert summary.records_skipped_already_present == 0
    assert summary.records_rejected_ineligible == 0
    assert len(summary.rows_appended_indices) == 184
    assert summary.rows_appended_indices[0] == 192
    assert summary.rows_appended_indices[-1] == 375

    # 2. Reopen and verify workbook integrity
    wb = openpyxl.load_workbook(out_file, data_only=False)
    ws = wb["تحليلي مشتريات"]

    # 3. Verify Table6.ref includes every data row (expanded to A2:L375)
    tbl = ws.tables["Table6"]
    assert tbl.ref == "A2:L375"

    # 4. Verify Grand-Total row remains below Table6 (at row 376) with summary formulas intact
    val_summary_title = ws.cell(376, 1).value
    assert val_summary_title == "الإجمالي"
    assert ws.cell(376, 6).value == "=SUBTOTAL(109,Table6[القيمة])"
    assert ws.cell(376, 7).value == "=SUBTOTAL(109,Table6[الضريبة])"
    assert ws.cell(376, 8).value == "=SUBTOTAL(109,Table6[الإجمالي])"

    # Verify merged cells shifted to row 376
    merged_str_list = [str(r) for r in ws.merged_cells.ranges]
    assert "A376:E376" in merged_str_list
    assert "I376:L376" in merged_str_list

    # 5. Verify the overflow row (row 375)
    rec_184 = records[-1]
    assert str(ws.cell(375, 2).value)[:10] == rec_184.date
    assert ws.cell(375, 3).value == rec_184.project
    assert ws.cell(375, 4).value == int(rec_184.invoice_number)
    assert ws.cell(375, 5).value == rec_184.supplier_name
    assert abs(float(ws.cell(375, 7).value) - rec_184.tax_amount) < 0.001

    # Verify Value and Total formulas in overflow row
    val_f375 = str(ws.cell(375, 6).value or "")
    tot_f375 = str(ws.cell(375, 8).value or "")
    assert val_f375.startswith("=") and "Table6" in val_f375
    assert tot_f375.startswith("=") and "Table6" in tot_f375

    # Verify styling and number formats in overflow row
    assert "Simplified Arabic" in (ws.cell(375, 3).font.name or "")
    assert ws.cell(375, 2).number_format == "yyyy-mm-dd"
    assert "0.00" in ws.cell(375, 7).number_format

    wb.close()
