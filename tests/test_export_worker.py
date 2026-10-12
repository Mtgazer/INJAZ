"""
Dedicated Unit and Integration Test Suite for ExportWorker Constructor and Dispatching.
Verifies:
1. Constructor signature consistency (keyword and positional arguments).
2. Proper dispatching to XlsxExporter and CsvPopulationEngine.
3. Neither format requiring the legacy Excel template.
4. Signal emission of export_finished only after file is validated on disk.
5. Signal emission of export_failed with Arabic message on error.
6. MainWindow._on_export_started and _on_csv_export_started integration without crashing.
"""
import os
import sys
import openpyxl
import pytest

os.environ["QT_QPA_PLATFORM"] = "offscreen"

from PySide6.QtWidgets import QApplication, QDialog, QMessageBox
from PySide6.QtCore import Qt

from injaz.models import (
    TransactionRecord,
    ExtractionStatus,
    InclusionEligibility,
    UserDecision,
    DuplicateStatus,
)
from injaz.ui.workers import ExportWorker, CsvExportWorker, XlsxExportWorker
from injaz.xlsx_exporter import XlsxExportSummary
from injaz.csv_writer import CsvExportSummary
from injaz.ui.main_window import MainWindow


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance()
    if app is None:
        app = QApplication(sys.argv)
    app.setLayoutDirection(Qt.RightToLeft)
    return app


@pytest.fixture(autouse=True)
def mock_dialogs(monkeypatch):
    """Mocks blocking dialog execution for headless testing."""
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.Yes)
    monkeypatch.setattr(QDialog, "exec", lambda self: QDialog.Accepted)
    monkeypatch.setattr(QDialog, "exec_", lambda self: QDialog.Accepted)


def make_record(index: int, invoice: str = "INV-100", tax: float = 150.0) -> TransactionRecord:
    return TransactionRecord(
        page_number=1,
        record_index=index,
        date="2026-06-15",
        transaction_type="قيد يدوي",
        raw_description=f"شراء مواد مشروع تجريبي فاتورة {invoice}",
        reference=f"REF-{index:04d}",
        debit=tax,
        credit=0.0,
        balance="15000.0",
        project="مشروع العاصمة",
        invoice_number=invoice,
        supplier_name="شركة المقاولات المتحدة",
        tax_amount=tax,
        extraction_status=ExtractionStatus.VALIDATED,
        inclusion_eligibility=InclusionEligibility.ELIGIBLE_FOR_INCLUSION,
        user_decision=UserDecision.NOT_APPLICABLE,
        duplicate_status=DuplicateStatus.UNIQUE,
    )


def test_export_worker_init_with_output_path_keyword(tmp_path):
    """Verifies ExportWorker accepts output_path and records as keyword arguments without TypeError."""
    out = str(tmp_path / "test.xlsx")
    records = [make_record(1)]

    worker = ExportWorker(
        output_path=out,
        records=records,
        export_format="xlsx",
        overwrite=True
    )
    assert worker.output_path == out
    assert worker.records == records
    assert worker.export_format == "xlsx"
    assert worker.overwrite is True


def test_export_worker_init_with_csv_format(tmp_path):
    """Verifies ExportWorker accepts export_format='csv'."""
    out = str(tmp_path / "test.csv")
    records = [make_record(1)]

    worker = ExportWorker(
        output_path=out,
        records=records,
        export_format="csv"
    )
    assert worker.output_path == out
    assert worker.records == records
    assert worker.export_format == "csv"


def test_export_worker_auto_detects_extension(tmp_path):
    """Verifies ExportWorker infers format from extension when appropriate."""
    csv_out = str(tmp_path / "auto.csv")
    w_csv = ExportWorker(output_path=csv_out, records=[make_record(1)])
    assert w_csv.export_format == "csv"

    xlsx_out = str(tmp_path / "auto.xlsx")
    w_xlsx = ExportWorker(output_path=xlsx_out, records=[make_record(1)])
    assert w_xlsx.export_format == "xlsx"


def test_export_worker_legacy_positional_signatures(tmp_path):
    """Verifies legacy positional calls (template, output, records) remain compatible."""
    records = [make_record(1)]
    tpl = str(tmp_path / "tpl.xlsx")
    out = str(tmp_path / "out.xlsx")

    # 3-arg legacy: (template_path, output_path, records)
    worker_3 = ExportWorker(tpl, out, records)
    assert worker_3.template_path == tpl
    assert worker_3.output_path == out
    assert worker_3.records == records
    assert worker_3.export_format == "xlsx"

    # 2-arg positional: (output_path, records)
    worker_2 = ExportWorker(out, records)
    assert worker_2.output_path == out
    assert worker_2.records == records
    assert worker_2.export_format == "xlsx"


def test_export_worker_xlsx_execution_end_to_end(qapp, tmp_path):
    """Verifies ExportWorker runs XLSX export, emits finished signal, and creates validated file."""
    out_xlsx = str(tmp_path / "worker_exported.xlsx")
    records = [make_record(1, "INV-1"), make_record(2, "INV-2")]

    worker = ExportWorker(
        output_path=out_xlsx,
        records=records,
        export_format="xlsx",
        overwrite=True
    )

    success_results = []
    error_results = []
    worker.export_finished.connect(lambda res: success_results.append(res))
    worker.export_failed.connect(lambda err: error_results.append(err))

    worker.run()

    assert len(error_results) == 0
    assert len(success_results) == 1
    summary = success_results[0]
    assert isinstance(summary, XlsxExportSummary)
    assert summary.records_exported == 2
    assert os.path.exists(out_xlsx)

    # Validate workbook contents
    wb = openpyxl.load_workbook(out_xlsx, data_only=True)
    ws = wb["المعاملات"]
    assert ws.max_row == 3  # 1 header + 2 data rows
    assert ws.cell(row=2, column=4).value == "INV-1"
    assert ws.cell(row=3, column=4).value == "INV-2"
    wb.close()


def test_export_worker_csv_execution_end_to_end(qapp, tmp_path):
    """Verifies ExportWorker runs CSV export, emits finished signal, and creates validated file."""
    out_csv = str(tmp_path / "worker_exported.csv")
    records = [make_record(1, "INV-1"), make_record(2, "INV-2")]

    worker = ExportWorker(
        output_path=out_csv,
        records=records,
        export_format="csv",
        overwrite=True
    )

    success_results = []
    error_results = []
    worker.export_finished.connect(lambda res: success_results.append(res))
    worker.export_failed.connect(lambda err: error_results.append(err))

    worker.run()

    assert len(error_results) == 0
    assert len(success_results) == 1
    summary = success_results[0]
    assert isinstance(summary, CsvExportSummary)
    assert summary.records_exported == 2
    assert os.path.exists(out_csv)

    # Validate CSV contents
    with open(out_csv, "r", encoding="utf-8-sig") as f:
        lines = [line.strip() for line in f if line.strip()]
    assert len(lines) == 3
    assert "INV-1" in lines[1]
    assert "INV-2" in lines[2]


def test_export_worker_error_propagation_arabic(qapp, tmp_path):
    """Verifies exceptions in export propagate via export_failed signal with Arabic prefix."""
    # Attempting to export with no eligible records triggers NoEligibleRecordsError
    worker = ExportWorker(
        output_path=str(tmp_path / "output.xlsx"),
        records=[],
        export_format="xlsx",
        overwrite=True
    )

    errors = []
    worker.export_failed.connect(lambda err: errors.append(err))
    worker.run()

    assert len(errors) == 1
    assert "خطأ أثناء تصدير ملف الإكسيل" in errors[0]


def test_main_window_export_started_handlers(qapp, tmp_path):
    """Verifies MainWindow._on_export_started and _on_csv_export_started run without TypeError."""
    win = MainWindow()
    out_xlsx = str(tmp_path / "main_win_out.xlsx")
    out_csv = str(tmp_path / "main_win_out.csv")
    records = [make_record(1)]

    # 1. Test _on_export_started (XLSX)
    win.output_path = out_xlsx
    win.records = records
    win._on_export_started(out_xlsx, records)

    assert win.export_worker is not None
    assert win.export_worker.output_path == out_xlsx
    assert win.export_worker.export_format == "xlsx"
    win.export_worker.wait(2000)

    # 2. Test _on_csv_export_started (CSV)
    win._on_csv_export_started(out_csv, records)
    assert win.csv_export_worker is not None
    assert win.csv_export_worker.output_path == out_csv
    assert win.csv_export_worker.export_format == "csv"
    win.csv_export_worker.wait(2000)

    win.close()
